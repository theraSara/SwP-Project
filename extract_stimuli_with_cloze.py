import re
import math
import pandas as pd
from docx import Document


rt = pd.read_csv("data/SPRT_LogLin_216.csv")
true_item_ids = sorted(rt.ITEM.unique())

doc = Document("data/Stimuli_Appendix_format.docx")
table = doc.tables[0]

rows = []
triplet_buffer = []
triplet_index = 0

EPSILON = 1e-6

for row in table.rows:
    cell = row.cells[0]

    raw_text = cell.text.strip().replace("*", "")
    if not raw_text:
        continue

    if raw_text.startswith(("H:", "M:", "L:")):
        critical_word_raw = ""
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                if run.underline:
                    critical_word_raw += run.text

        critical_word = critical_word_raw.strip().strip(".,;!?'\"")

        condition_letter = raw_text[0]
        condition_map = {"H": "HC", "M": "MC", "L": "LC"}
        condition = condition_map[condition_letter]

        sentence_part = raw_text[2:].strip()
        cloze_match = re.search(r"\((\d+(?:\.\d+)?)%?\)\s*$", sentence_part)

        if cloze_match:
            cloze_percent = float(cloze_match.group(1))
            cloze_prob = cloze_percent / 100.0

            zero_cloze = (cloze_prob == 0.0)

            if cloze_prob > 0:
                cloze_surprisal_raw = -math.log2(cloze_prob)
            else:
                cloze_surprisal_raw = float("inf")

            cloze_surprisal_smoothed = -math.log2(max(cloze_prob, EPSILON))

            sentence_clean = re.sub(
                r"\((\d+(?:\.\d+)?)%?\)\s*$",
                "",
                sentence_part
            ).strip()

        else:
            cloze_percent = None
            cloze_prob = None
            zero_cloze = None
            cloze_surprisal_raw = None
            cloze_surprisal_smoothed = None
            sentence_clean = sentence_part.strip()

        if critical_word in sentence_clean:
            parts = sentence_clean.rsplit(critical_word, 1)
            prefix = parts[0].rstrip()
            suffix = parts[1].lstrip()
        else:
            prefix = "ERROR"
            suffix = "ERROR"

        target_llm = " " + critical_word

        triplet_buffer.append({
            "condition": condition,
            "sentence": sentence_clean,
            "prefix": prefix,
            "suffix": suffix,
            "target_llm": target_llm,
            "critical_word": critical_word,

            "cloze_percent": cloze_percent,
            "cloze_prob": cloze_prob,
            "zero_cloze": zero_cloze,
            "cloze_surprisal_raw": cloze_surprisal_raw,
            "cloze_surprisal_smoothed": cloze_surprisal_smoothed,
        })

        if len(triplet_buffer) == 3:
            true_item = true_item_ids[triplet_index]
            triplet_index += 1

            for item in triplet_buffer:
                item["ITEM"] = true_item
                rows.append(item)

            triplet_buffer = []

df = pd.DataFrame(rows)

print("Rows:", len(df))
print("Unique ITEM:", df.ITEM.nunique())

errors = df[df["prefix"] == "ERROR"]
if len(errors) > 0:
    print(f"Warning: {len(errors)} rows failed to match the critical word.")

missing_cloze = df["cloze_prob"].isna().sum()
if missing_cloze > 0:
    print(f"Warning: {missing_cloze} rows have no cloze probability extracted.")

zero_count = (df["zero_cloze"] == True).sum()
print(f"Zero-cloze items: {zero_count}")

df.to_csv("data_output/bk21_stimuli_with_cloze.csv", index=False)
print("Stimuli extracted successfully!")